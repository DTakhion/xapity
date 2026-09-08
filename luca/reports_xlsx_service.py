# luca/reports_xlsx_service.py

"""
Generación determinista de reportes XLSX de Luca.

Responsabilidades:

- recibir un snapshot de reporte previamente validado;
- construir un archivo Excel en memoria;
- preservar la estructura contable del reporte.

Este módulo NO:

- consulta MongoDB;
- resuelve usuarios;
- envía correos;
- modifica acciones pendientes o confirmadas.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import (
    get_column_letter,
)

from luca.reports_delivery_service import (
    ConfirmedReportSnapshot,
)


GENERAL_BALANCE_COLUMNS = (
    (
        "numCuenta",
        "Código",
    ),
    (
        "nameCuenta",
        "Cuenta",
    ),
    (
        "debe",
        "Debe",
    ),
    (
        "haber",
        "Haber",
    ),
    (
        "deudor",
        "Deudor",
    ),
    (
        "acreedor",
        "Acreedor",
    ),
    (
        "activo",
        "Activo",
    ),
    (
        "pasivo",
        "Pasivo",
    ),
    (
        "perdida",
        "Pérdida",
    ),
    (
        "ganancia",
        "Ganancia",
    ),
)


@dataclass(
    frozen=True,
    slots=True,
)
class GeneratedReportFile:
    """
    Archivo de reporte generado en memoria.
    """

    filename: str
    content_type: str
    content: bytes

    @property
    def size_bytes(
        self,
    ) -> int:
        return len(
            self.content
        )


def _write_report_row(
    *,
    worksheet,
    row_number: int,
    row_data: dict,
    bold: bool = False,
) -> None:
    """
    Escribe una fila utilizando la estructura
    oficial del Balance General.
    """

    for column_number, (
        field_name,
        _,
    ) in enumerate(
        GENERAL_BALANCE_COLUMNS,
        start=1,
    ):
        value = row_data.get(
            field_name
        )

        cell = worksheet.cell(
            row=row_number,
            column=column_number,
            value=value,
        )

        if bold:
            cell.font = Font(
                bold=True
            )

        if column_number >= 3:
            cell.number_format = (
                '#,##0;[Red]-#,##0'
            )

            cell.alignment = Alignment(
                horizontal="right"
            )


def generate_general_balance_xlsx(
    *,
    snapshot: ConfirmedReportSnapshot,
) -> GeneratedReportFile:
    """
    Genera un XLSX a partir de un Balance General
    previamente validado.

    El archivo es construido completamente en memoria.
    """

    if not isinstance(
        snapshot,
        ConfirmedReportSnapshot,
    ):
        raise TypeError(
            "snapshot debe ser una instancia de "
            "ConfirmedReportSnapshot."
        )

    report = snapshot.report

    accounts = (
        report.get(
            "accounts"
        )
        or []
    )

    if not isinstance(
        accounts,
        list,
    ):
        raise TypeError(
            "report.accounts debe ser una lista."
        )

    subtotal = (
        report.get(
            "subtotal"
        )
    )

    result = (
        report.get(
            "result"
        )
    )

    totals = (
        report.get(
            "totals"
        )
    )

    for field_name, value in (
        (
            "subtotal",
            subtotal,
        ),
        (
            "result",
            result,
        ),
        (
            "totals",
            totals,
        ),
    ):
        if not isinstance(
            value,
            dict,
        ):
            raise TypeError(
                f"report.{field_name} debe ser "
                "un diccionario."
            )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Balance General"
    )

    # ==============================================
    # ESTILOS
    # ==============================================

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    subtotal_fill = PatternFill(
        fill_type="solid",
        fgColor="EDEDED",
    )

    result_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAD3",
    )

    thin_gray = Side(
        style="thin",
        color="D9D9D9",
    )

    border = Border(
        bottom=thin_gray,
    )

    # ==============================================
    # TÍTULO
    # ==============================================

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(
            GENERAL_BALANCE_COLUMNS
        ),
    )

    title_cell = worksheet.cell(
        row=1,
        column=1,
        value="Balance General",
    )

    title_cell.font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )

    title_cell.fill = (
        title_fill
    )

    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[
        1
    ].height = 26

    # ==============================================
    # METADATOS
    # ==============================================

    worksheet.cell(
        row=3,
        column=1,
        value="Período",
    ).font = Font(
        bold=True
    )

    worksheet.cell(
        row=3,
        column=2,
        value=(
            f"{snapshot.query_from} "
            f"a {snapshot.query_to}"
        ),
    )

    worksheet.cell(
        row=4,
        column=1,
        value="Versión",
    ).font = Font(
        bold=True
    )

    worksheet.cell(
        row=4,
        column=2,
        value=snapshot.version,
    )

    # ==============================================
    # ENCABEZADOS
    # ==============================================

    header_row = 6

    for column_number, (
        _,
        label,
    ) in enumerate(
        GENERAL_BALANCE_COLUMNS,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_number,
            value=label,
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = (
            header_fill
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = (
            border
        )

    # ==============================================
    # CUENTAS
    # ==============================================

    current_row = (
        header_row + 1
    )

    for account in accounts:
        if not isinstance(
            account,
            dict,
        ):
            raise TypeError(
                "Cada elemento de accounts "
                "debe ser un diccionario."
            )

        _write_report_row(
            worksheet=worksheet,
            row_number=current_row,
            row_data=account,
        )

        current_row += 1

    # ==============================================
    # SUBTOTAL
    # ==============================================

    _write_report_row(
        worksheet=worksheet,
        row_number=current_row,
        row_data=subtotal,
        bold=True,
    )

    for cell in worksheet[
        current_row
    ]:
        cell.fill = (
            subtotal_fill
        )

    current_row += 1

    # ==============================================
    # RESULTADO
    # ==============================================

    _write_report_row(
        worksheet=worksheet,
        row_number=current_row,
        row_data=result,
        bold=True,
    )

    for cell in worksheet[
        current_row
    ]:
        cell.fill = (
            result_fill
        )

    current_row += 1

    # ==============================================
    # TOTALES
    # ==============================================

    _write_report_row(
        worksheet=worksheet,
        row_number=current_row,
        row_data=totals,
        bold=True,
    )

    for cell in worksheet[
        current_row
    ]:
        cell.fill = (
            total_fill
        )

    # ==============================================
    # PRESENTACIÓN
    # ==============================================

    worksheet.freeze_panes = (
        "A7"
    )

    worksheet.auto_filter.ref = (
        f"A6:"
        f"{get_column_letter(len(GENERAL_BALANCE_COLUMNS))}"
        f"{header_row + len(accounts)}"
    )

    column_widths = {
        1: 16,
        2: 38,
        3: 18,
        4: 18,
        5: 18,
        6: 18,
        7: 18,
        8: 18,
        9: 18,
        10: 18,
    }

    for (
        column_number,
        width,
    ) in column_widths.items():
        worksheet.column_dimensions[
            get_column_letter(
                column_number
            )
        ].width = width

    worksheet.sheet_view.showGridLines = (
        False
    )

    # ==============================================
    # METADATA TÉCNICA
    # ==============================================

    metadata_sheet = (
        workbook.create_sheet(
            "Metadata"
        )
    )

    metadata_rows = (
        (
            "businessId",
            snapshot.business_id,
        ),
        (
            "reportType",
            snapshot.report_type,
        ),
        (
            "queryFrom",
            snapshot.query_from,
        ),
        (
            "queryTo",
            snapshot.query_to,
        ),
        (
            "mongoId",
            snapshot.mongo_id,
        ),
        (
            "version",
            snapshot.version,
        ),
        (
            "contentHash",
            snapshot.content_hash,
        ),
    )

    for row_number, (
        key,
        value,
    ) in enumerate(
        metadata_rows,
        start=1,
    ):
        metadata_sheet.cell(
            row=row_number,
            column=1,
            value=key,
        ).font = Font(
            bold=True
        )

        metadata_sheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

    metadata_sheet.column_dimensions[
        "A"
    ].width = 20

    metadata_sheet.column_dimensions[
        "B"
    ].width = 72

    # ==============================================
    # SERIALIZACIÓN
    # ==============================================

    buffer = BytesIO()

    workbook.save(
        buffer
    )

    content = (
        buffer.getvalue()
    )

    filename = (
        "balance_general_"
        f"{snapshot.query_from}_"
        f"{snapshot.query_to}.xlsx"
    )

    return GeneratedReportFile(
        filename=filename,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        content=content,
    )